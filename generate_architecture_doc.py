import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left=None, right=None):
    tcPr = cell._element.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    borders_xml += f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>' if top else '<w:top w:val="none"/>'
    borders_xml += f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>' if bottom else '<w:bottom w:val="none"/>'
    borders_xml += f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{left}"/>' if left else '<w:left w:val="none"/>'
    borders_xml += f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{right}"/>' if right else '<w:right w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def build_architecture_doc(output_path):
    doc = docx.Document()

    # Set page margins (0.75 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Style colors
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy #0F172A
    COLOR_SECONDARY = RGBColor(14, 165, 233)  # Cyber Blue #0EA5E9
    COLOR_DARK = RGBColor(30, 41, 59)         # Slate Dark #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B

    # Helper functions
    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK
        return p

    def add_diagram_image(image_path):
        if os.path.exists(image_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run()
            run.add_picture(image_path, width=Inches(7.0))
            
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_before = Pt(2)
            cap.paragraph_format.space_after = Pt(14)
            c_run = cap.add_run("Figure 1: Provigil AETHER End-to-End Enterprise Architecture, Ingress Routing & Multi-Agent Topology")
            c_run.font.name = "Calibri"
            c_run.font.size = Pt(9.5)
            c_run.font.italic = True
            c_run.font.color.rgb = COLOR_MUTED

    # ---------------- DOCUMENT HEADER ----------------
    add_title("PROVIGIL AETHER - XDR PLATFORM")
    add_subtitle("Comprehensive System Architecture & Multi-Agent Topology Specification")
    
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    r = meta_p.add_run("Classification: ")
    r.bold = True
    r.font.size = Pt(9.5)
    meta_p.add_run("Proprietary & Confidential  |  ").font.size = Pt(9.5)
    r2 = meta_p.add_run("Platform Version: ")
    r2.bold = True
    r2.font.size = Pt(9.5)
    meta_p.add_run("v2.4 (Unified NDR + SIEM)  |  ").font.size = Pt(9.5)
    r3 = meta_p.add_run("Date: ")
    r3.bold = True
    r3.font.size = Pt(9.5)
    meta_p.add_run("October 2026").font.size = Pt(9.5)

    # ---------------- SECTION 1: EXECUTIVE SUMMARY ----------------
    add_h1("1. Executive Summary & Architectural Overview")
    add_body("The Provigil AETHER platform is a unified Extended Detection and Response (XDR) enterprise cybersecurity solution that harmonizes Network Detection & Response (NDR) with Security Information and Event Management (SIEM). Engineered in high-performance Rust and modern Angular 19, the architecture provides zero-compromise multi-tenancy, processing tens of thousands of high-velocity telemetry events per second with sub-second threat detection and analytical query capabilities.")
    add_body("The ecosystem is divided into four distinct operational tiers: Ingress Edge Gateway (Caddy), Microservices Application Tier (Rust Auth and SIEM API), Ingestion & Agent Fleet (Wazuh Windows/Linux Endpoint Agents and Network Sensors), and the Columnar Big Data & Caching Tier (ClickHouse & Valkey).")

    # ---------------- SECTION 2: END-TO-END SYSTEM ARCHITECTURE DIAGRAM ----------------
    add_h1("2. End-to-End System Architecture Diagram")
    add_body("The high-resolution diagram below illustrates the comprehensive topology including edge TLS ingress, reverse proxy URI dispatching, Rust backend microservices, ClickHouse big data columnar storage, Wazuh remoted wire protocols, and distributed endpoint agents:")

    diagram_img = "Provigil_AETHER_Complete_Architecture_Diagram.png"
    add_diagram_image(diagram_img)


    # ---------------- SECTION 3: COMPONENT DETAIL SPECIFICATIONS ----------------
    add_h1("3. Detailed Core Components & Daemon Subsystems Breakdown")

    add_h2("3.1 Edge Reverse Proxy (Caddy Alpine)")
    add_bullet("The single unified edge ingress gateway running in Docker container xdr-caddy-proxy on ports 80 (HTTP) and 443 (HTTPS/QUIC).", "Role: ")
    add_bullet("Terminates TLS with automated certificate renewal, enforces strict security headers (nosniff, SAMEORIGIN, strict-origin-when-cross-origin), and enables zstd/gzip compression.", "Security & Performance: ")
    add_bullet("Inspects request URI paths and distributes traffic dynamically without requiring client port awareness:", "Smart Upstream Routing: ")
    add_bullet("Forwards to auth-service:3001.", "  • /api/auth/* & /api/v1/auth/*: ")
    add_bullet("Forwards to siem-api:8088.", "  • /api/*, /api/v1/*, & /downloads/*: ")
    add_bullet("Forwards full-duplex WebSocket connections to siem-api:8088.", "  • /ws/alerts & /ws/siem/*: ")
    add_bullet("Serves the compiled production Angular 19 SPA with HTML5 push-state fallback.", "  • Root /*: ")

    add_h2("3.2 Authentication & Licensing Service (auth-service)")
    add_bullet("Compiled native binary built in Rust using Tokio, Axum, and SQLx/ClickHouse.", "Technology: ")
    add_bullet("Issues and verifies HMAC-SHA256 JWT tokens containing user identity, active tenant_id, role scope, and licensed features.", "Authentication: ")
    add_bullet("Validates whether a tenant has access to 'ndr', 'siem', or unified 'xdr' modules, dynamically enforcing UI rendering and route authorization.", "Feature Entitlement: ")
    add_bullet("Utilizes Valkey on port 6379 for millisecond token revocation checks, rate limiting, and session verification.", "Session Caching: ")

    add_h2("3.3 Core SIEM API & Processing Engine (siem-api)")
    add_bullet("Rust-based multi-threaded event processing engine listening on HTTP/REST (8088), Syslog UDP (514), and Syslog TCP (601).", "Architecture: ")
    add_bullet("Executes over 1,400+ Wazuh and Sigma XML/YAML rules across Windows EventLog, Linux Auditd, and network streams.", "Detection Engine: ")
    add_bullet("Maintains an in-memory temporal sliding window to detect brute force, lateral movement, port scans, and distributed attacks.", "Sliding Window Correlation: ")
    add_bullet("Broadcasts matching alerts in real time to connected Analyst SOC browsers via /ws/alerts.", "WebSocket Push Bus: ")
    add_bullet("Handles POST /api/v1/agents/enroll requests, authenticating tenants via X-Tenant-Key and generating client.keys entries.", "Agent Enrollment Handler: ")

    add_h2("3.4 Core Manager Daemons & Wire Subsystems")
    add_bullet("Listens on Port 1514 (TCP/UDP) with low-overhead AES-256 wire encryption. Capable of sustaining 50,000+ simultaneous persistent agent connections per cluster node with mutual client.keys cryptographic authentication.", "wazuh-remoted (Agent Wire Gateway): ")
    add_bullet("High-speed pre-decoding and XML/Sigma rule evaluation daemon. Correlates multi-event security incident chains and outputs normalized siem_alerts.", "wazuh-analysisd (Detection Engine): ")
    add_bullet("SQLite database daemon managing agent synchronization, hardware inventories, and persistent host state.", "wazuh-db (Agent State Synchronization): ")
    add_bullet("Continuously evaluates inventory discovered on host agents against the National Vulnerability Database (NVD) and OVAL CVE feeds to flag vulnerable software.", "wazuh-vuln-detector (CVE Audit Engine): ")
    add_bullet("Orchestrates alert dispatches and automated incident ticketing to Slack, VirusTotal, Shuffle, Jira, and Email notifications.", "siem-integratord & siem-maild: ")

    # ---------------- SECTION 4: COMPLETE HOST AGENT SUBSYSTEMS ----------------
    add_h1("4. Comprehensive Endpoint Agent Architecture & Host Internal Subsystems")
    add_body("The Provigil endpoint agent fleet runs lightweight native collectors embedded directly into host operating systems. The internal architecture of each agent is divided into dedicated, modular subsystems designed for sub-1% CPU consumption and zero business disruption:")

    add_h2("4.1 Windows Host Internal Subsystems (WazuhSvc / wazuh-agent.exe)")
    add_bullet("Operates as a registered Windows Service under %ProgramFiles(x86)%\\ossec-agent\\wazuh-agent.exe.", "Host Installation: ")
    add_bullet("Streams raw Windows EventChannel events in real time across Application, Security, System, PowerShell, DHCP, and IIS channels.", "Logcollector Engine: ")
    add_bullet("Deep kernel-level telemetry covering Event IDs 1-20: Process Creation & Command-Line arguments (ID 1), Network Connections & Remote IPs (ID 3), Driver Loads (ID 6), DLL Injection / Image Loads (ID 7), CreateRemoteThread / Mimikatz injection (ID 8), ProcessAccess (ID 10), and Registry Modifications (IDs 12-14).", "Sysmon Telemetry Integration: ")
    add_bullet("Periodically audits the complete host profile: CPU/RAM hardware specs, OS build & hotfixes, network adapters, MAC addresses, IPv4/IPv6 routing tables, open/listening network ports, running processes & PIDs, and installed software registry packages.", "Syscollector (Host Inventory): ")
    add_bullet("Continuous cryptographic hashing (SHA256/SHA1/MD5) of critical system directories and Windows Registry persistence keys (Run, RunOnce, Winlogon, Services, AppInit_DLLs). Features real-time Whodata tracking that identifies the exact Windows user account responsible for file alterations.", "File & Registry Integrity Monitoring (FIM): ")
    add_bullet("Automated system hardening audits evaluated against Center for Internet Security (CIS) benchmarks, BitLocker status, Windows Defender real-time protection, and password policy compliance.", "SCA & Rootcheck: ")
    add_bullet("Executes local automated containment commands dispatched by the SIEM: dynamic Windows Firewall IP drops, malicious process termination (taskkill), and compromised user account locking.", "Active Response (execd): ")
    add_bullet("Features an in-memory leaky-bucket buffer backed by an encrypted disk queue, guaranteeing zero telemetry loss during network disruptions.", "Agent Resilience Queue: ")

    add_h2("4.2 Linux Host Internal Subsystems (siem-agent-linux)")
    add_bullet("Native lightweight daemon managed under systemd as a high-priority system service.", "Host Architecture: ")
    add_bullet("Intercepts kernel-level system calls via Netlink socket, tracking execve process launches, socket connect/bind operations, ptrace debugging/injection, and privileged sudo executions.", "Linux Auditd Kernel Hook: ")
    add_bullet("Continuously streams systemd journald logs, /var/log/secure, auth.log, SSH brute force events, and application logs (Nginx, Apache, MySQL, PostgreSQL).", "Journald & Logcollector: ")
    add_bullet("Kernel inotify filesystem watchers monitoring /etc, /bin, /sbin, /usr/bin, package managers (dpkg, rpm, pacman), and /proc process tree states.", "Linux Syscollector & FIM: ")
    add_bullet("Hardening checks against CIS Linux Distribution benchmarks, SSH daemon security options, sysctl kernel parameters, and stealth rootkit detection.", "SCA & Rootcheck: ")
    add_bullet("Instantly drops attacker IP addresses via iptables / nftables and terminates compromised shell processes in under 1 second.", "Active Response: ")

    add_h2("4.3 Network Sensor & Appliance Subsystems (NDR Module)")
    add_bullet("Connects to physical or virtual network TAP/SPAN ports to passively inspect Ethernet traffic without impacting throughput.", "Inspection Model: ")
    add_bullet("Extracts protocol metadata, JA4/JA4S TLS cryptographic fingerprints, unencrypted HTTP payloads, and DNS tunneling anomalies.", "Zeek & Suricata Analyzers: ")
    add_bullet("Monitors ARP broadcast tables to detect ARP poisoning, rogue DHCP servers, and gateway spoofing in real time.", "ARP Anomaly Engine (ndr-arp.zeek): ")
    add_bullet("Receives standard RFC 5424/6587 syslog streams from FortiGate, Palo Alto, Cisco FTD, Juniper, and enterprise routers.", "Syslog Appliance Ingestion: ")
    add_bullet("Uses SSH and SNMP protocols via siem-agentlessd to monitor routers, switches, and firewalls where software agents cannot be installed.", "Agentless Monitoring (siem-agentlessd): ")


    # ---------------- SECTION 5: NETWORK COMMUNICATION & PORT MATRIX ----------------
    add_h1("5. Network Communication & Port Matrix")
    add_body("The following matrix summarizes all network ports, transport protocols, and communication flows across the platform:")

    headers = ["Port", "Protocol", "Direction", "Source", "Destination", "Purpose & Security"]
    data = [
        ["80 / 443", "TCP/QUIC", "Inbound", "Web Browsers", "Caddy Proxy", "HTTPS Web UI, REST API, & WebSockets (TLS)"],
        ["3001", "TCP", "Internal", "Caddy Proxy", "auth-service", "JWT Auth, Tenant/User provisioning, RBAC"],
        ["8088", "TCP", "Inbound/Int", "Caddy / Agents", "siem-api", "REST Endpoints, /downloads, Agent Enrollment"],
        ["1514", "TCP/UDP", "Inbound", "Endpoint Agents", "remoted / SIEM", "Encrypted Agent Telemetry & Keepalives (AES)"],
        ["1515", "TCP", "Internal", "Agent CLI / Ops", "wazuh-authd", "Legacy agent registration daemon"],
        ["514 / 601", "TCP/UDP", "Inbound", "Network Devices", "siem-api", "Syslog RFC 5424 / RFC 6587 Ingestion"],
        ["8123", "TCP", "Internal", "Microservices", "ClickHouse", "HTTP Query & Batch Ingestion Port"],
        ["9000", "TCP", "Internal", "CLI / Native", "ClickHouse", "High-speed Native TCP Client Protocol"],
        ["6379", "TCP", "Internal", "Microservices", "Valkey / Redis", "In-memory token caching, session pub/sub"],
    ]

    col_widths = [Inches(0.9), Inches(0.9), Inches(0.9), Inches(1.3), Inches(1.3), Inches(1.7)]
    
    table = doc.add_table(rows=len(data) + 1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Format data rows
    for r_idx, row_values in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            set_cell_borders(row_cells[c_idx], top="E2E8F0", bottom="E2E8F0")
            p = row_cells[c_idx].paragraphs[0]
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(9.0)
            run.font.color.rgb = COLOR_DARK
            if c_idx == 0:
                run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------- SECTION 6: MULTI-TENANCY & DATA ISOLATION ----------------
    add_h1("6. Multi-Tenancy & Cryptographic Isolation Model")
    add_body("Provigil AETHER implements a multi-layered tenant isolation architecture ensuring complete segregation of data, credentials, and management controls across customer organizations:")
    add_bullet("Every agent installation command embeds a unique 64-character tenant key (e.g., tk_187e9a96...). The enrollment API binds the newly spawned agent permanently to that specific tenant in ClickHouse (ndr.tenant_agent_keys and ndr.agent_registry).", "Agent Provisioning Isolation: ")
    add_bullet("All core telemetry tables in ClickHouse (siem_events, siem_alerts, siem_agents, xdr_incidents) are partitioned and sorted by tenant_id. Queries executed by analysts automatically inject WHERE tenant_id = :scope, preventing data leakage across tenant boundaries.", "Storage Partitioning: ")
    add_bullet("Super Admins can govern global system health and tenant licenses; Tenant Admins can provision analysts, toggle SIEM/NDR suites, and manage local agent enrollments; Analysts are constrained to alert triage and incident investigation for their assigned tenant only.", "Role-Based Access Control (RBAC): ")

    # ---------------- SECTION 7: STEP-BY-STEP AGENT ENROLLMENT LIFECYCLE ----------------
    add_h1("7. Agent Deployment & Enrollment Lifecycle")
    add_body("Understanding the precise handshake flow eliminates connectivity issues and explains the transition from PENDING to ACTIVE:")
    add_bullet("Administrator copies the one-line deployment script from the Tenant Admin or Fleet UI. The script specifies the SIEM Manager URL (http://<SIEM_IP>:8088) and the X-Tenant-Key.", "1. Script Generation: ")
    add_bullet("PowerShell downloads wazuh-agent.exe, configuration files, and dynamically queries /downloads/windows/ossec.conf?manager=<IP> to write an IPv4-targeted ossec.conf.", "2. Package Retrieval: ")
    add_bullet("The script makes a POST request to /api/v1/agents/enroll. The SIEM API validates the tenant key, issues a sequential agent ID (e.g. 019), and generates a 256-bit AES cryptographic key saved into client.keys.", "3. Registration: ")
    add_bullet("The script installs and starts the 'WazuhSvc' service. The agent initiates an AES handshake on TCP port 1514 to the manager.", "4. Service Initialization: ")
    add_bullet("Upon receiving the first keepalive message and syscollector evaluation over port 1514, the SIEM engine updates the database, transitioning the agent status from PENDING to ACTIVE.", "5. State Activation: ")

    # Save document
    doc.save(output_path)
    print(f"Architecture document successfully generated at: {output_path}")

if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "Provigil_AETHER_System_Architecture.docx"
    build_architecture_doc(out)
